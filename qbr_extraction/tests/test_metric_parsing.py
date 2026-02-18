from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from qbr_intelligence.metrics.parsing import parse_pptx_deck


def _build_fixture(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    textbox = slide.shapes.add_textbox(100, 100, 300, 50)
    textbox.text = "CTR: 2.5%\nSpend $1.2M"

    table_shape = slide.shapes.add_table(2, 2, 100, 200, 300, 100)
    table = table_shape.table
    table.cell(0, 0).text = "Reach"
    table.cell(0, 1).text = "100K"
    table.cell(1, 0).text = "Unique Reach"
    table.cell(1, 1).text = "80K"

    notes = slide.notes_slide
    notes.notes_text_frame.text = "Speaker notes: Impressions 1.2M"

    prs.save(path)


def test_parse_pptx_deck(tmp_path: Path) -> None:
    deck_path = tmp_path / "fixture.pptx"
    _build_fixture(deck_path)

    deck = parse_pptx_deck(deck_path)
    assert deck.deck_id == "fixture"
    assert len(deck.slides) == 1

    slide = deck.slides[0]
    assert slide.slide_index == 1
    assert any("CTR" in block.text for block in slide.text_blocks)
    assert len(slide.tables) == 1
    assert any(cell.text == "Reach" for cell in slide.tables[0].cells)
    assert slide.notes_blocks
    assert "Impressions" in slide.notes_blocks[0].text
