"""Backfill slide titles in DB using PPTX slide title extraction."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from sqlalchemy import create_engine, text


def _to_sync_database_url(database_url: str) -> str:
    return database_url.replace("+aiosqlite", "").replace("+asyncpg", "")


def _iter_shapes(shapes):
    for shape in shapes:
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _normalize_candidate(text: str) -> str:
    if not text:
        return ""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for line in lines:
        cleaned = re.sub(r"^[-*]\s+", "", line).strip()
        if not cleaned:
            continue
        if cleaned.startswith("![") or cleaned.startswith("<"):
            continue
        if cleaned.lower().startswith("source:"):
            continue
        return re.sub(r"\s+", " ", cleaned)[:500]
    return ""


def _shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    frame = shape.text_frame
    text = (frame.text or "").strip()
    if text:
        return text
    chunks: list[str] = []
    for para in frame.paragraphs:
        part = "".join(run.text for run in para.runs).strip() if para.runs else (para.text or "").strip()
        if part:
            chunks.append(part)
    return "\n".join(chunks).strip()


def _shape_font_size_points(shape) -> float:
    if not getattr(shape, "has_text_frame", False):
        return 0.0
    sizes: list[float] = []
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font is not None and run.font.size is not None:
                sizes.append(float(run.font.size.pt))
    if not sizes:
        return 0.0
    return max(sizes)


def _is_title_placeholder(shape) -> bool:
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        kind = shape.placeholder_format.type
    except Exception:
        return False
    return kind in {
        PP_PLACEHOLDER.TITLE,
        PP_PLACEHOLDER.CENTER_TITLE,
        PP_PLACEHOLDER.SUBTITLE,
    }


def _extract_slide_title(slide) -> str:
    main_title_shape = getattr(slide.shapes, "title", None)
    if main_title_shape is not None:
        title = _normalize_candidate(_shape_text(main_title_shape))
        if title:
            return title

    best_score: float | None = None
    best_title = ""
    for shape in _iter_shapes(slide.shapes):
        text_value = _normalize_candidate(_shape_text(shape))
        if not text_value:
            continue
        try:
            top = int(shape.top)
        except Exception:
            top = 0
        score = 0.0
        if _is_title_placeholder(shape):
            score += 10000.0
        score += _shape_font_size_points(shape) * 100.0
        score += max(0, 600000 - top) / 100.0
        if len(text_value) <= 100:
            score += 40.0
        if best_score is None or score > best_score:
            best_score = score
            best_title = text_value
    return best_title


def extract_titles_from_pptx(pptx_path: str | Path) -> dict[int, str]:
    prs = Presentation(str(pptx_path))
    out: dict[int, str] = {}
    for idx, slide in enumerate(prs.slides, start=1):
        title = _extract_slide_title(slide).strip()
        if title:
            out[idx] = title
    return out


def title_from_raw_text(raw_text: str | None) -> str:
    text = (raw_text or "").replace("\r\n", "\n")
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("![") or line.startswith("<"):
            continue
        if line.lower().startswith("### notes:"):
            continue
        if line.lower().startswith("source:"):
            continue
        line = re.sub(r"^[-*]\s+", "", line).strip()
        if line:
            return re.sub(r"\s+", " ", line)[:500]
    return ""


@dataclass(frozen=True)
class SlideTitleBackfillStats:
    document_id: int
    updated: int
    skipped_existing: int
    skipped_empty: int
    skipped_same: int
    source_pptx_used: bool


def backfill_slide_titles_for_document(
    *,
    database_url: str,
    document_id: int,
    pptx_path: str | Path | None = None,
    overwrite: bool = False,
    dry_run: bool = False,
) -> SlideTitleBackfillStats:
    sync_url = _to_sync_database_url(database_url)
    engine = create_engine(sync_url, future=True)
    titles_from_pptx: dict[int, str] = {}
    source_pptx_used = False

    if pptx_path:
        pptx = Path(pptx_path)
        if pptx.exists():
            try:
                titles_from_pptx = extract_titles_from_pptx(pptx)
                source_pptx_used = bool(titles_from_pptx)
            except Exception:
                titles_from_pptx = {}

    updated = 0
    skipped_existing = 0
    skipped_empty = 0
    skipped_same = 0

    with engine.begin() as conn:
        rows = conn.execute(
            text(
                """
                SELECT slide_number, title, raw_text
                FROM slides
                WHERE document_id = :document_id
                ORDER BY slide_number ASC
                """
            ),
            {"document_id": document_id},
        ).mappings().all()

        for row in rows:
            slide_number = int(row["slide_number"])
            current_title = (row["title"] or "").strip()
            new_title = (titles_from_pptx.get(slide_number) or "").strip()
            if not new_title:
                new_title = title_from_raw_text(row["raw_text"])
            if not new_title:
                skipped_empty += 1
                continue
            if current_title and not overwrite:
                skipped_existing += 1
                continue
            if current_title == new_title:
                skipped_same += 1
                continue
            updated += 1
            if not dry_run:
                conn.execute(
                    text(
                        """
                        UPDATE slides
                        SET title = :title
                        WHERE document_id = :document_id
                          AND slide_number = :slide_number
                        """
                    ),
                    {
                        "title": new_title,
                        "document_id": document_id,
                        "slide_number": slide_number,
                    },
                )

    engine.dispose()
    return SlideTitleBackfillStats(
        document_id=document_id,
        updated=updated,
        skipped_existing=skipped_existing,
        skipped_empty=skipped_empty,
        skipped_same=skipped_same,
        source_pptx_used=source_pptx_used,
    )
