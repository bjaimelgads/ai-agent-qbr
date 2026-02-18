"""Adapters for parsing PPTX into canonical slide models."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Iterable
import json
import re
import zipfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from qbr_intelligence.metrics.models import (
    BBox,
    DeckContent,
    NotesBlock,
    SlideContent,
    Table,
    TableCell,
    TextBlock,
)


def _iter_shapes(shapes) -> Iterable:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _shape_bbox(shape) -> BBox | None:
    try:
        return BBox(
            left=int(shape.left),
            top=int(shape.top),
            width=int(shape.width),
            height=int(shape.height),
        )
    except Exception:
        return None


def _extract_text_from_shape(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    text_frame = shape.text_frame
    lines: list[str] = []
    for para in text_frame.paragraphs:
        if para.runs:
            text = "".join(run.text for run in para.runs)
        else:
            text = para.text
        if text and text.strip():
            lines.append(text.strip())
    if not lines:
        return (text_frame.text or "").strip()
    return "\n".join(lines)


def _extract_notes(slide, slide_index: int) -> list[NotesBlock]:
    try:
        if not slide.has_notes_slide:
            return []
        notes_slide = slide.notes_slide
        notes_frame = notes_slide.notes_text_frame
        notes_text = (notes_frame.text or "").strip()
        if not notes_text:
            return []
        return [
            NotesBlock(
                slide_index=slide_index,
                block_id=f"{slide_index}:notes",
                text=notes_text,
            )
        ]
    except Exception:
        return []


_TABLE_RE = re.compile(r"<table>(?P<body>.*?)</table>", re.IGNORECASE | re.DOTALL)
_ROW_RE = re.compile(r"<tr>(?P<body>.*?)</tr>", re.IGNORECASE | re.DOTALL)
_CELL_RE = re.compile(r"<t[dh]>(?P<body>.*?)</t[dh]>", re.IGNORECASE | re.DOTALL)


def parse_pptx_deck(path: str | Path) -> DeckContent:
    deck_path = Path(path)
    try:
        prs = Presentation(str(deck_path))
    except (zipfile.BadZipFile, KeyError, ValueError):
        fallback = _parse_from_extraction_output(deck_path)
        if fallback is not None:
            return fallback
        raise
    slides: list[SlideContent] = []
    for index, slide in enumerate(prs.slides, start=1):
        text_blocks: list[TextBlock] = []
        tables: list[Table] = []
        for shape in _iter_shapes(slide.shapes):
            shape_id = getattr(shape, "shape_id", None)
            if shape_id is None:
                continue
            if getattr(shape, "has_table", False):
                table = shape.table
                table_id = f"{index}:{shape_id}"
                cells: list[TableCell] = []
                for r, row in enumerate(table.rows):
                    for c, _ in enumerate(row.cells):
                        cell = table.cell(r, c)
                        text = (cell.text or "").strip()
                        if not text:
                            continue
                        cell_id = f"{table_id}:r{r}c{c}"
                        cells.append(
                            TableCell(
                                slide_index=index,
                                table_id=table_id,
                                row=r,
                                col=c,
                                text=text,
                                cell_id=cell_id,
                            )
                        )
                tables.append(
                    Table(
                        slide_index=index,
                        table_id=table_id,
                        cells=tuple(cells),
                        nrows=len(table.rows),
                        ncols=len(table.columns),
                    )
                )
                continue
            if getattr(shape, "has_text_frame", False):
                text = _extract_text_from_shape(shape)
                if text:
                    block_id = f"{index}:{shape_id}"
                    text_blocks.append(
                        TextBlock(
                            slide_index=index,
                            block_id=block_id,
                            text=text,
                            source_type="slide_text",
                            bbox=_shape_bbox(shape),
                        )
                    )
        notes_blocks = _extract_notes(slide, index)
        slides.append(
            SlideContent(
                slide_index=index,
                text_blocks=tuple(text_blocks),
                tables=tuple(tables),
                notes_blocks=tuple(notes_blocks),
            )
        )
    deck_id = deck_path.stem
    return DeckContent(deck_id=deck_id, slides=tuple(slides))


def _parse_from_extraction_output(deck_path: Path) -> DeckContent | None:
    repo_root = Path(__file__).resolve().parents[3]
    candidates = [
        repo_root / "qbr_extraction" / "qbr_pipeline" / "output",
    ]
    deck_name = deck_path.stem
    for base in candidates:
        if not base.exists():
            continue
        target = base / deck_name
        if not target.exists():
            underscored = base / deck_name.replace(" ", "_")
            if underscored.exists():
                target = underscored
            else:
                continue
        slides_path = target / "03_slides_parsed.json"
        if not slides_path.exists():
            continue
        payload = json.loads(slides_path.read_text(encoding="utf-8"))
        slides: list[SlideContent] = []
        for slide in payload:
            slide_index = int(slide.get("slide_number", 0))
            raw_text = (slide.get("raw_text") or "").strip()
            notes_text = slide.get("speaker_notes") or ""
            text_blocks: list[TextBlock] = []
            tables: list[Table] = []
            notes_blocks: list[NotesBlock] = []

            if raw_text:
                tables, raw_text = _extract_tables_from_html(raw_text, slide_index)
                text_blocks.append(
                    TextBlock(
                        slide_index=slide_index,
                        block_id=f"{slide_index}:raw_text",
                        text=raw_text,
                        source_type="slide_text",
                        bbox=None,
                    )
                )
            if notes_text:
                notes_blocks.append(
                    NotesBlock(
                        slide_index=slide_index,
                        block_id=f"{slide_index}:notes",
                        text=notes_text,
                    )
                )
            slides.append(
                SlideContent(
                    slide_index=slide_index,
                    text_blocks=tuple(text_blocks),
                    tables=tuple(tables),
                    notes_blocks=tuple(notes_blocks),
                )
            )
        return DeckContent(deck_id=deck_name, slides=tuple(slides))
    return None


def _extract_tables_from_html(text: str, slide_index: int) -> tuple[list[Table], str]:
    tables: list[Table] = []
    cleaned = text
    for table_idx, table_match in enumerate(_TABLE_RE.finditer(text), start=1):
        body = table_match.group("body")
        rows = []
        for row_match in _ROW_RE.finditer(body):
            row_body = row_match.group("body")
            cells = [cell_match.group("body").strip() for cell_match in _CELL_RE.finditer(row_body)]
            if cells:
                rows.append(cells)
        if not rows:
            continue
        table_id = f"{slide_index}:html_table_{table_idx}"
        cells: list[TableCell] = []
        for r, row in enumerate(rows):
            for c, cell_text in enumerate(row):
                if not cell_text:
                    continue
                cell_id = f"{table_id}:r{r}c{c}"
                cells.append(
                    TableCell(
                        slide_index=slide_index,
                        table_id=table_id,
                        row=r,
                        col=c,
                        text=cell_text,
                        cell_id=cell_id,
                    )
                )
        tables.append(
            Table(
                slide_index=slide_index,
                table_id=table_id,
                cells=tuple(cells),
                nrows=len(rows),
                ncols=max(len(row) for row in rows) if rows else 0,
            )
        )
    if tables:
        cleaned = _TABLE_RE.sub("", cleaned).strip()
        cleaned = cleaned.replace("### Notes:", "").strip()
    return tables, cleaned


def compute_deck_hash(path: str | Path) -> str:
    deck_path = Path(path)
    digest = hashlib.sha256()
    with deck_path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
